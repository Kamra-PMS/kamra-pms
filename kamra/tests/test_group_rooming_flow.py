"""The desk's group flow: block rooms, name a rooming list, book the named
guests now and keep the rest of the block held for the group."""

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import add_days, nowdate

from kamra.api import (add_occupant_names, available_rooms, create_group_with_rooming_list,
                       group_detail, rooming_list_template)


class GroupRoomingFlowTestCase(IntegrationTestCase):
	def setUp(self):
		frappe.set_user("Administrator")  # nosemgrep: frappe-setuser -- test fixture
		frappe.local.lang = "en"
		self.property = frappe.get_doc({"doctype": "Property",
			"property_name": "TEST GRF " + frappe.generate_hash(length=8),
			"country": "Uganda", "currency": "UGX"}).insert(ignore_permissions=True).name
		self.room_type = frappe.get_doc({"doctype": "Room Type", "property": self.property,
			"room_type_code": "STD", "room_type_name": "Standard", "base_price": 1000,
			"adults_capacity": 2}).insert(ignore_permissions=True).name
		for n in ("101", "102", "103", "104", "105"):
			frappe.get_doc({"doctype": "Room", "property": self.property,
				"room_type": self.room_type, "room_number": n}).insert(ignore_permissions=True)
		self.ci = str(add_days(nowdate(), 10))
		self.co = str(add_days(nowdate(), 12))

	def make(self, blocked=4, guests=None, **kw):
		return create_group_with_rooming_list(
			self.property, "Sharma Wedding", self.ci, self.co,
			[{"room_type": self.room_type, "rooms_blocked": blocked}],
			guests if guests is not None else [
				{"room_type": self.room_type, "guest_name": "Anil Sharma", "phone": "+256700000011"},
				{"room_type": self.room_type, "guest_name": "Priya Sharma"},
			],
			phone="+256700000011", **kw)

	def test_named_guests_are_booked_and_the_rest_stay_held(self):
		out = self.make(master_index=0, master_pays="Room only")
		self.assertEqual(out["skipped"], [])
		self.assertEqual((out["blocked"], len(out["created"]), out["held"]), (4, 2, 2))
		gb = frappe.get_doc("Group Booking", out["group_booking"])
		self.assertEqual(gb.status, "Confirmed")
		self.assertEqual(gb.master_reservation, out["created"][0])
		detail = group_detail(gb.name)
		self.assertEqual(len(detail["rooming_list"]), 2)
		pick = next(p for p in detail["pickup"] if p["room_type"] == self.room_type)
		self.assertEqual((pick["picked_up"], pick["remaining"]), (2, 2))
		# the 2 held rooms are off general sale: 5 rooms - 2 booked - 2 held
		self.assertEqual(len(available_rooms(self.property, self.room_type, self.ci, self.co)), 1)

	def test_a_block_bigger_than_the_free_rooms_leaves_nothing_behind(self):
		before = frappe.db.count("Group Booking", {"property": self.property})
		with self.assertRaises(frappe.ValidationError):
			self.make(blocked=6)
		self.assertEqual(frappe.db.count("Group Booking", {"property": self.property}), before)

	def test_a_past_arrival_is_refused_before_the_group_is_made(self):
		before = frappe.db.count("Group Booking", {"property": self.property})
		with self.assertRaises(frappe.ValidationError):
			create_group_with_rooming_list(
				self.property, "Late", str(add_days(nowdate(), -1)), self.co,
				[{"room_type": self.room_type, "rooms_blocked": 1}], [])
		self.assertEqual(frappe.db.count("Group Booking", {"property": self.property}), before)

	def test_more_guests_than_blocked_rooms_is_refused(self):
		with self.assertRaises(frappe.ValidationError):
			self.make(blocked=1)

	def test_an_empty_rooming_list_only_blocks(self):
		out = self.make(guests=[])
		self.assertEqual((out["created"], out["held"]), ([], 4))

	def test_others_staying_go_on_the_room_register(self):
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
		                         "others": "Ravi Sharma; Meera Sharma; anil sharma"}])
		res = frappe.get_doc("Reservation", out["created"][0])
		# the room's own guest is never listed twice
		self.assertEqual([o.full_name for o in res.occupants], ["Ravi Sharma", "Meera Sharma"])

	def test_add_occupant_names_keeps_names_already_there(self):
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma"}])
		res = out["created"][0]
		self.assertEqual(add_occupant_names(res, "Ravi Sharma, Meera Sharma")["added"], 2)
		self.assertEqual(add_occupant_names(res, ["Meera Sharma", "Kiran"])["added"], 1)
		self.assertEqual(len(frappe.get_doc("Reservation", res).occupants), 3)

	def test_precheckin_fills_the_register(self):
		from kamra.public_api import _merge_precheckin_occupants
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
		                         "others": "Ravi Sharma"}])
		res = frappe.get_doc("Reservation", out["created"][0])
		_merge_precheckin_occupants(res, [
			{"full_name": "Ravi Sharma", "age": "34", "id_type": "Passport", "id_number": "P123"},
			{"full_name": "Little Sharma", "age": "6", "id_type": "Passport", "id_number": "KID1"},
			{"full_name": "", "age": "40"},
		])
		occ = {o.full_name: o for o in frappe.get_doc("Reservation", res.name).occupants}
		self.assertEqual(set(occ), {"Ravi Sharma", "Little Sharma"})
		# an age from the link is only what the guest says, for the desk to confirm
		self.assertEqual(occ["Ravi Sharma"].id_number, "P123")
		self.assertEqual((occ["Ravi Sharma"].age, occ["Ravi Sharma"].guest_stated_age), (0, 34))
		self.assertEqual((occ["Little Sharma"].age, occ["Little Sharma"].guest_stated_age), (0, 6))

	def test_a_name_with_no_age_is_an_adult_until_told_otherwise(self):
		from kamra.checkin import checkin_flow
		from kamra.public_api import _merge_precheckin_occupants, precheckin_info
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
		                         "others": "Ravi Sharma"}])
		res = frappe.get_doc("Reservation", out["created"][0])
		from kamra.api import registration_card
		self.assertIsNone(registration_card(res.name)["occupants"][0]["age"])
		# the link shows an empty age, not 0
		info = precheckin_info(res.precheckin_token)
		self.assertIsNone(info["occupants"][0]["age"])
		# and the desk still asks for this person's ID
		row = checkin_flow(res.name)["ids"]["occupants"][0]
		self.assertEqual((row["age"], row["exempt"]), (None, False))
		# an ID typed on the link is kept, even if the form sends 0 for the age
		_merge_precheckin_occupants(res, [{"row": info["occupants"][0]["row"],
			"full_name": "Ravi Sharma", "age": "0", "id_type": "Passport", "id_number": "P9"}])
		saved = frappe.get_doc("Reservation", res.name).occupants[0]
		self.assertEqual((saved.id_number, saved.guest_stated_age), ("P9", 0))

	def _with_billed_child(self):
		"""A stay with a child being charged, and a free-child age of 6."""
		frappe.db.set_value("Room Type", self.room_type, {
			"free_child_age": 6, "adults_capacity": 4, "children_capacity": 2,
			"child_price": 500, "child_age_limit": 12})
		frappe.clear_document_cache("Room Type", self.room_type)
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
		                         "adults": 2, "children": 1}])
		return frappe.get_doc("Reservation", out["created"][0])

	def test_an_age_from_the_guest_link_never_changes_the_price(self):
		from kamra.public_api import _merge_precheckin_occupants
		from kamra.pricing import quote
		res = self._with_billed_child()
		total = res.amount_after_tax
		key = res.rate_lock_key
		# the child really is being charged, so a free age would lower this
		no_child = quote(self.property, self.room_type, self.ci, self.co, 2, 0)
		self.assertGreater(total, no_child["amount_after_tax"])
		_merge_precheckin_occupants(res, [{"full_name": "Little Sharma", "age": "2"}])
		res = frappe.get_doc("Reservation", res.name)
		self.assertEqual((res.amount_after_tax, res.rate_lock_key), (total, key))
		# nor does the next ordinary save at the desk
		res.special_requests = "Late arrival"
		res.save()
		self.assertEqual(frappe.get_doc("Reservation", res.name).amount_after_tax, total)

	def test_the_guest_link_can_rename_and_remove_but_not_touch_a_confirmed_row(self):
		from kamra.public_api import _merge_precheckin_occupants, precheckin_info
		out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
		                         "others": "Ravi Sharma; Meera Sharma; Desk Guest"}])
		res = frappe.get_doc("Reservation", out["created"][0])
		rows = {o.full_name: o for o in res.occupants}
		frappe.db.set_value("Stay Occupant", rows["Desk Guest"].name, "age", 40)
		_merge_precheckin_occupants(res, [
			{"row": rows["Ravi Sharma"].name, "full_name": "Ravi K Sharma"},
			{"row": rows["Desk Guest"].name, "full_name": "Changed By Guest"},
		])
		names = sorted(o.full_name for o in frappe.get_doc("Reservation", res.name).occupants)
		# renamed in place (no duplicate), Meera removed, the confirmed row untouched
		self.assertEqual(names, ["Desk Guest", "Ravi K Sharma"])

	def test_a_room_is_priced_for_the_people_named_in_it(self):
		out = self.make(guests=[
			{"room_type": self.room_type, "guest_name": "Solo"},
			{"room_type": self.room_type, "guest_name": "Pair", "others": "Partner"},
			{"room_type": self.room_type, "guest_name": "Crowd", "others": "One; Two; Three"},
		])
		adults = [frappe.db.get_value("Reservation", r, "adults") for r in out["created"]]
		self.assertEqual(adults, [1, 2, 2])  # capped at what the room sleeps

	def test_a_shared_phone_keeps_each_typed_name(self):
		out = self.make(guests=[
			{"room_type": self.room_type, "guest_name": "Anil Sharma", "phone": "+256700000099"},
			{"room_type": self.room_type, "guest_name": "Priya Sharma", "phone": "+256700000099"},
		])
		names = [frappe.db.get_value("Reservation", r, "guest_name") for r in out["created"]]
		self.assertEqual(names, ["Anil Sharma", "Priya Sharma"])

	def test_names_that_fail_to_save_are_a_warning_not_a_skipped_room(self):
		from unittest.mock import patch
		with patch("kamra.api.add_occupant_names", side_effect=frappe.ValidationError("boom")):
			out = self.make(guests=[{"room_type": self.room_type, "guest_name": "Anil Sharma",
			                         "others": "Ravi"}])
		self.assertEqual((len(out["created"]), out["skipped"]), (1, []))
		self.assertEqual(out["warnings"][0]["reservation"], out["created"][0])

	def test_rooming_list_template_lists_the_room_types(self):
		from io import BytesIO

		from openpyxl import load_workbook
		rooming_list_template(self.property)
		wb = load_workbook(BytesIO(frappe.response["filecontent"]))
		ws = wb["Rooming List"]
		self.assertEqual([c.value for c in ws[1]], ["Name", "Phone", "Room type", "Other guests"])
		self.assertEqual(wb["Lists"]["A1"].value, "Standard")
		self.assertEqual(wb["Lists"].sheet_state, "hidden")
		self.assertIn("Lists!", ws.data_validations.dataValidation[0].formula1)
